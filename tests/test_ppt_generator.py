"""PPT 生成工具测试（假数据，不调 LLM）。"""

import pytest
from pptx import Presentation

from app.tools.ppt_generator import deck_to_pptx

FAKE_DECK = {
    "outline": {
        "title": "用例图建模入门",
        "subtitle": "从需求到用例图的四步迭代",
        "sections": [],
    },
    "pages": [
        {
            "title": "课程目标",
            "bullets": ["理解用例图三要素", "掌握边界画法", "能识别常见错误"],
            "note": "讲解要点：三要素概念辨析。时间：约 5 分钟。",
            "minutes": 5,
        },
        {
            "title": "什么是参与者",
            "bullets": ["外部角色", "画在边界外"],
            "note": "结合教材 AI互动④ 三查追问。",
            "minutes": 3,
        },
    ],
}


def test_deck_to_pptx_creates_file(tmp_path):
    """生成 .pptx 且页数正确（封面 + 内容页）。"""
    out = deck_to_pptx(FAKE_DECK, tmp_path / "out.pptx")
    assert out.exists() and out.stat().st_size > 1024
    prs = Presentation(str(out))
    assert len(prs.slides) == 3  # 封面 + 2 页


def test_deck_to_pptx_notes_attached(tmp_path):
    """讲课备注应写入 notes（PRD PPT-4）。"""
    out = deck_to_pptx(FAKE_DECK, tmp_path / "out.pptx")
    prs = Presentation(str(out))
    note = prs.slides[1].notes_slide.notes_text_frame.text
    assert "三要素" in note


def test_deck_to_pptx_rejects_empty(tmp_path):
    """空页面 deck 应拒绝生成（避免交付空文件）。"""
    with pytest.raises(ValueError):
        deck_to_pptx({"outline": {"title": "x"}, "pages": []}, tmp_path / "empty.pptx")
