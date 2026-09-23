"""实训任务工单生成链：课题 → 结构化工单 JSON → docx 导出。

复用 app.core.structured_output（JSON 提取/校验/重试/Fake 注入），
响应契约对齐前端 renderTasksheetResults（task_id/course_name/…）。
"""

import asyncio

from app.config.settings import get_settings
from app.core.structured_output import invoke_structured
from app.models.tasksheet import Tasksheet

_SYSTEM = (
    "你是高职院校软件技术专业的实训教学设计专家，负责生成规范的实训任务工单。\n"
    "工单结构要求（产教融合范式）：\n"
    "1. task_id 格式为 WS-2026-<课程缩写><两位序号>，如 WS-2026-SE01；\n"
    "2. occupational_role 使用企业岗位名（如：软件建模师、系统分析师）；\n"
    "3. training_mode 使用企业化训练方式（如：双人结对协作、项目组接力）；\n"
    "4. scenario 是一段 100~300 字的任务情景描述，模拟企业真实项目背景；\n"
    "5. deliverables 是 3~4 项成果交付物（每项含 name/format/desc）；\n"
    "6. steps 是 4~6 道阶梯工序（每道含 step_num/step_name/"
    "estimated_time（纯数字分钟）/guide 操作指引/checkpoint 关键检验点）；\n"
    "7. 严格只输出一个 JSON 对象，不要输出任何解释文字。"
)

_USER_TMPL = (
    "请为以下实训课题生成任务工单。\n"
    "课题：{topic}\n"
    "课时：{minutes} 分钟\n"
    "难度：{difficulty}\n"
    "训练模式：{mode}\n\n"
    "输出 JSON 格式（严格遵循）：\n"
    "{{\n"
    '  "task_id": "WS-2026-SE01",\n'
    '  "course_name": "课题名",\n'
    '  "occupational_role": "岗位名",\n'
    '  "training_mode": "训练方式",\n'
    '  "difficulty": "难度",\n'
    '  "duration_minutes": 45,\n'
    '  "scenario": "任务情景…",\n'
    '  "deliverables": [{{"name": "成果名", "format": "格式", "desc": "说明"}}],\n'
    '  "steps": [{{"step_num": 1, "step_name": "工序名", "estimated_time": "15", '
    '"guide": "操作指引", "checkpoint": "检验点"}}]\n'
    "}}\n"
    "要求：duration_minutes 为数字；estimated_time 为纯数字分钟字符串。"
)


def _generate_sync(topic: str, minutes: int, difficulty: str, mode: str, llm_factory=None) -> Tasksheet:
    messages = [
        ("system", _SYSTEM),
        ("user", _USER_TMPL.format(topic=topic, minutes=minutes, difficulty=difficulty, mode=mode)),
    ]
    return invoke_structured(
        messages,
        Tasksheet,
        llm_factory=llm_factory,
        retries=1,
        timeout=get_settings().gen_llm_timeout,
    )


async def generate_tasksheet(topic: str, minutes: int, difficulty: str, mode: str, *, llm_factory=None) -> Tasksheet:
    """异步生成实训任务工单（测试可注入 llm_factory）。"""
    return await asyncio.to_thread(_generate_sync, topic, minutes, difficulty, mode, llm_factory)


def export_tasksheet_docx(ts: Tasksheet) -> bytes:
    """工单导出为 docx 字节流（python-docx，供下载）。"""
    import io

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(f"实训任务工单：{ts.course_name}", level=1)
    meta = doc.add_paragraph()
    meta.add_run(f"工单编号：{ts.task_id}　岗位角色：{ts.occupational_role}　").bold = True
    meta.add_run(f"训练模式：{ts.training_mode}　难度：{ts.difficulty}　时长：{ts.duration_minutes} 分钟").bold = True

    doc.add_heading("一、任务情景", level=2)
    doc.add_paragraph(ts.scenario)

    doc.add_heading("二、成果交付物", level=2)
    for d in ts.deliverables:
        doc.add_paragraph(f"{d.name}（{d.format}）：{d.desc}", style="List Bullet")

    doc.add_heading("三、阶梯工序", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["工序", "名称", "时长(分钟)", "操作指引", "关键检验点"]):
        cell = table.rows[0].cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
    for s in ts.steps:
        row = table.add_row().cells
        row[0].text = str(s.step_num)
        row[1].text = s.step_name
        row[2].text = s.estimated_time
        row[3].text = s.guide
        row[4].text = s.checkpoint
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)

    doc.add_heading("四、提交要求", level=2)
    doc.add_paragraph("按工序顺序完成并自检关键检验点，全部交付物打包提交，文件名含工单编号。")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
