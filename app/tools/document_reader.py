"""文档解析器：把教材源稿（md/docx/pdf）解析为带标题元数据的文本块。

首批入库的是教材源稿 Markdown 提取稿（`02-需求文档PRD/_教材源稿提取/`），
docx/pdf 解析能力为后续教材扩展预留。
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

# 行首为 "# 标题" 或 "任务X" 的作为新块的边界
_HEADING_RE = re.compile(r"^(#{1,4})\s+(.+)$")
_TASK_RE = re.compile(r"^(任务[一二三四五六七八九十]+[^\n]*)$")


@dataclass
class DocumentChunk:
    """一个带来源与标题的文本块。"""

    text: str
    source: str
    heading: str = ""


@dataclass
class ParsedDocument:
    """一份文档的解析结果。"""

    path: str
    chunks: list[DocumentChunk] = field(default_factory=list)


def _push_chunk(
    chunks: list[DocumentChunk],
    source: str,
    heading: str,
    buf: list[str],
) -> None:
    """把缓冲行组成一个非空块并加入列表。"""
    text = "\n".join(buf).strip()
    if text:
        chunks.append(DocumentChunk(text=text, source=source, heading=heading))


def parse_markdown(filepath: str | Path, source_name: str | None = None) -> ParsedDocument:
    """按标题层级把 Markdown 切块，保留标题作为元数据。

    Args:
        filepath: md 文件路径。
        source_name: 来源名（如"教材·智能问答"）；缺省用文件名。

    Returns:
        ParsedDocument: 含块列表；块 text 保留章节内容，heading 为最近标题。
    """
    path = Path(filepath)
    source = source_name or path.stem
    chunks: list[DocumentChunk] = []
    current_heading = ""
    buf: list[str] = []

    for line in path.read_text(encoding="utf-8").splitlines():
        m = _TASK_RE.match(line)
        if m:  # 教材"任务X"行（无 # 前缀）也作为新块边界
            _push_chunk(chunks, source, current_heading, buf)
            buf = [line]
            current_heading = line.strip()
            continue
        h = _HEADING_RE.match(line)
        if h:  # 普通 md 标题
            _push_chunk(chunks, source, current_heading, buf)
            buf = [line]
            current_heading = h.group(2).strip()
            continue
        buf.append(line)
    _push_chunk(chunks, source, current_heading, buf)

    return ParsedDocument(path=str(path), chunks=chunks)


def parse_markdown_dir(
    directory: str | Path,
    suffix: str = "*.md",
    source_prefix: str = "",
) -> list[DocumentChunk]:
    """解析目录下全部 md 文件，返回合并后的块列表。

    Args:
        directory: 教材源稿目录。
        suffix: 文件后缀匹配。
        source_prefix: 来源名前缀（如"教材"）。

    Returns:
        list[DocumentChunk]: 合并的全部块。
    """
    all_chunks: list[DocumentChunk] = []
    for filepath in sorted(Path(directory).glob(suffix)):
        doc = parse_markdown(filepath, source_name=f"{source_prefix}{filepath.stem}")
        all_chunks.extend(doc.chunks)
    return all_chunks


def read_docx_text(filepath: str | Path) -> str:
    """读取 docx 全文（后续教材扩展用）。"""
    import docx  # 延迟导入，减少无该依赖时的启动成本

    document = docx.Document(str(filepath))
    return "\n".join(p.text for p in document.paragraphs)


def read_pdf_text(filepath: str | Path) -> str:
    """读取 pdf 全文（后续教材扩展用）。"""
    from pypdf import PdfReader  # 延迟导入

    reader = PdfReader(str(filepath))
    return "\n".join(page.extract_text() or "" for page in reader.pages)
